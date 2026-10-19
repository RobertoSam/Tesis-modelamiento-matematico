# Genera la plantilla reference.docx con el formato institucional provisional:
# Arial 12, interlineado 1,5, márgenes 3 cm (izq.) y 2,5 cm (resto), número de página centrado abajo.
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

entrada, salida = sys.argv[1], sys.argv[2]
doc = Document(entrada)

def fijar_fuente(estilo, tam=None, negrita=None):
    estilo.font.name = "Arial"
    rpr = estilo.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), "Arial")
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]
    if tam: estilo.font.size = Pt(tam)
    if negrita is not None: estilo.font.bold = negrita
    estilo.font.color.rgb = RGBColor(0, 0, 0)

for nombre in ["Normal", "Body Text", "First Paragraph", "Compact", "Bibliography"]:
    try:
        e = doc.styles[nombre]
    except KeyError:
        continue
    fijar_fuente(e, 12)
    pf = e.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if nombre == "Bibliography":
        pf.left_indent = Cm(1.27); pf.first_line_indent = Cm(-1.27)

for nombre, tam in [("Title", 14), ("Subtitle", 12), ("Heading 1", 14), ("Heading 2", 12),
                    ("Heading 3", 12), ("Table Caption", 12), ("Image Caption", 12)]:
    try:
        e = doc.styles[nombre]
    except KeyError:
        continue
    fijar_fuente(e, tam, None if "Caption" in nombre else True)
    e.paragraph_format.space_before = Pt(12)
    e.paragraph_format.space_after = Pt(6)

for s in doc.sections:
    s.left_margin = Cm(3); s.right_margin = Cm(2.5)
    s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5)
    p = s.footer.paragraphs[0] if s.footer.paragraphs else s.footer.add_paragraph()
    p.alignment = 1
    run = p.add_run()
    for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipo:
            f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), tipo); run._r.append(f)
        else:
            t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = texto; run._r.append(t)

doc.save(salida)
print("plantilla generada:", salida)
